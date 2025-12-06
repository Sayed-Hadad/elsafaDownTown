from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, current_app
from flask_login import login_required, current_user
from app import db
from app.models import Booking, Hall, StaffLog, Addon
from datetime import datetime, timedelta
import csv
from io import StringIO

bookings_bp = Blueprint('bookings', __name__, url_prefix='/bookings')

@bookings_bp.route('/', methods=['GET'])
@login_required
def list_bookings():
    """List bookings with filters and search"""
    try:
        page = request.args.get('page', 1, type=int)
        if page < 1:
            page = 1
        search = request.args.get('search', '').strip()
        hall_id = request.args.get('hall_id', type=int)
        status = request.args.get('status', '').strip()
        
        query = Booking.query
        
        # Filter by staff member if not admin
        if not current_user.is_admin():
            query = query.filter_by(created_by=current_user.id)
        
        # Search by customer name or phone number
        if search:
            from sqlalchemy import or_
            query = query.filter(
                or_(
                    Booking.customer_name.ilike(f'%{search}%'),
                    Booking.phone.ilike(f'%{search}%'),
                    Booking.phone_alt.ilike(f'%{search}%')
                )
            )
        
        # Filter by hall
        if hall_id:
            query = query.filter_by(hall_id=hall_id)
        
        # Filter by status
        if status:
            query = query.filter_by(status=status)
        
        # Order and paginate
        bookings = query.order_by(Booking.date.desc(), Booking.start_time.desc()).paginate(
            page=page, 
            per_page=10,
            error_out=False
        )
        halls = Hall.query.filter_by(status='active').all()
        
        return render_template('bookings/list.html',
                             bookings=bookings,
                             halls=halls,
                             search=search,
                             hall_id=hall_id,
                             status=status)
    except Exception as e:
        flash(f'حدث خطأ في تحميل الحجوزات: {str(e)}', 'error')
        bookings = Booking.query.filter_by(created_by=current_user.id).paginate(page=1, per_page=10, error_out=False) if not current_user.is_admin() else Booking.query.paginate(page=1, per_page=10, error_out=False)
        halls = Hall.query.all()
        return render_template('bookings/list.html',
                             bookings=bookings,
                             halls=halls,
                             search='',
                             hall_id=None,
                             status='')

@bookings_bp.route('/add', methods=['GET', 'POST'])
@login_required
def add_booking():
    """Add new booking"""
    if request.method == 'POST':
        hall_id = request.form.get('hall_id', type=int)
        customer_name = request.form.get('customer_name', '').strip()
        phone = request.form.get('phone', '').strip()
        phone_alt = request.form.get('phone_alt', '').strip()
        date_str = request.form.get('date')
        # بالنسبة للحجز اليومي، نستخدم وقت تشغيل القاعة بالكامل
        hall = Hall.query.get(hall_id)
        if not hall:
            flash('القاعة غير موجودة', 'error')
            return redirect(url_for('bookings.add_booking'))

        start_time = hall.working_from
        end_time = hall.working_to
        deposit = request.form.get('deposit', type=float, default=0)
        total_price = request.form.get('total_price', type=float)
        notes = request.form.get('notes', '').strip()
        
        # Validate inputs
        if not all([hall_id, customer_name, phone, date_str, total_price is not None]):
            flash('الرجاء ملء جميع الحقول المطلوبة', 'error')
            return redirect(url_for('bookings.add_booking'))
        
        try:
            date = datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            flash('تاريخ غير صحيح', 'error')
            return redirect(url_for('bookings.add_booking'))
        
        # Check for conflicts
        conflicts = Booking.query.filter(
            Booking.hall_id == hall_id,
            Booking.date == date,
            Booking.status != 'cancelled'
        ).all()
        
        if conflicts:
            flash('اليوم محجوز بالكامل لهذه القاعة. الرجاء اختيار تاريخ آخر', 'error')
            return redirect(url_for('bookings.add_booking'))
        
        # Parse selected addons from DB
        selected = []
        selected_ids = request.form.getlist('addon_selected[]')
        selected_ids = [int(aid) for aid in selected_ids if aid]
        addons_lookup = {a.id: a for a in Addon.query.filter(Addon.id.in_(selected_ids)).all()}
        for addon_id in selected_ids:
            addon_model = addons_lookup.get(addon_id)
            if addon_model:
                addon = addon_model.to_dict()
                price_key = f'addon_price_{addon_id}'
                try:
                    addon['price'] = float(request.form.get(price_key, addon['price']))
                except (ValueError, TypeError):
                    pass
                note_key = f'addon_note_{addon_id}'
                addon['note'] = request.form.get(note_key, '')
                selected.append(addon)

        booking = Booking(
            hall_id=hall_id,
            customer_name=customer_name,
            phone=phone,
            phone_alt=phone_alt,
            date=date,
            start_time=start_time,
            end_time=end_time,
            deposit=deposit,
            total_price=total_price,
            notes=notes,
            addons=selected if selected else None,
            created_by=current_user.id,
            status='pending'
        )
        
        db.session.add(booking)
        db.session.commit()
        
        # Log action
        log = StaffLog(
            user_id=current_user.id,
            action='create_booking',
            details=f'حجز جديد: {customer_name} - {date}'
        )
        db.session.add(log)
        db.session.commit()
        
        flash(f'تم إضافة الحجز بنجاح - {customer_name}', 'success')
        return redirect(url_for('bookings.list_bookings'))
    
    halls = Hall.query.filter_by(status='active').all()
    # Prefill hall/date if provided via query params (from calendar)
    selected_hall = None
    prefill_date = request.args.get('date')
    hall_qid = request.args.get('hall_id', type=int)
    if hall_qid:
        selected_hall = Hall.query.get(hall_qid)

    addons = Addon.query.filter_by(active=True).order_by(Addon.name).all()
    addons_payload = [a.to_dict() for a in addons]

    return render_template('bookings/add.html',
                           halls=halls,
                           selected_hall=selected_hall,
                           prefill_date=prefill_date,
                           available_addons=addons_payload)

@bookings_bp.route('/<int:booking_id>/edit', methods=['GET', 'POST'])
@login_required
def edit_booking(booking_id):
    """Edit booking"""
    booking = Booking.query.get_or_404(booking_id)
    
    # Check permission
    if not current_user.is_admin() and booking.created_by != current_user.id:
        flash('ليس لديك صلاحيات لتعديل هذا الحجز', 'error')
        return redirect(url_for('bookings.list_bookings'))
    
    if request.method == 'POST':
        customer_name = request.form.get('customer_name', '').strip()
        phone = request.form.get('phone', '').strip()
        phone_alt = request.form.get('phone_alt', '').strip()
        date_str = request.form.get('date')
        start_time = booking.hall.working_from
        end_time = booking.hall.working_to
        deposit = request.form.get('deposit', type=float, default=0)
        total_price = request.form.get('total_price', type=float)
        status = request.form.get('status', booking.status)
        notes = request.form.get('notes', '').strip()
        
        if not all([customer_name, phone, date_str, total_price is not None]):
            flash('الرجاء ملء جميع الحقول المطلوبة', 'error')
            return redirect(url_for('bookings.edit_booking', booking_id=booking_id))
        
        try:
            date = datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            flash('تاريخ غير صحيح', 'error')
            return redirect(url_for('bookings.edit_booking', booking_id=booking_id))
        
        # Check for conflicts (excluding current booking)
        conflicts = Booking.query.filter(
            Booking.hall_id == booking.hall_id,
            Booking.date == date,
            Booking.id != booking_id,
            Booking.status != 'cancelled'
        ).all()
        
        if conflicts:
            flash('اليوم محجوز بالكامل لهذه القاعة. الرجاء اختيار تاريخ آخر', 'error')
            return redirect(url_for('bookings.edit_booking', booking_id=booking_id))
        
        # Parse DB add-ons
        selected_addon_ids = request.form.getlist('addon_selected')
        selected_addon_ids = [int(aid) for aid in selected_addon_ids if aid]
        addons_lookup = {a.id: a for a in Addon.query.filter(Addon.id.in_(selected_addon_ids)).all()}
        selected_addons = []

        for addon_id in selected_addon_ids:
            addon_model = addons_lookup.get(addon_id)
            if addon_model:
                addon_price = request.form.get(f'addon_price_{addon_id}', type=float, default=addon_model.price)
                addon_note = request.form.get(f'addon_note_{addon_id}', '').strip()
                selected_addons.append({
                    'id': addon_id,
                    'name': addon_model.name,
                    'price': addon_price,
                    'note': addon_note,
                    'icon': addon_model.icon or ''
                })
        
        booking.customer_name = customer_name
        booking.phone = phone
        booking.phone_alt = phone_alt
        booking.date = date
        booking.start_time = start_time
        booking.end_time = end_time
        booking.deposit = deposit
        booking.total_price = total_price
        booking.notes = notes
        booking.status = status
        booking.addons = selected_addons
        
        db.session.commit()
        
        # Log action
        log = StaffLog(
            user_id=current_user.id,
            action='edit_booking',
            details=f'تعديل الحجز: {customer_name} - {date}'
        )
        db.session.add(log)
        db.session.commit()
        
        flash(f'تم تحديث الحجز بنجاح', 'success')
        return redirect(url_for('bookings.list_bookings'))
    
    halls = Hall.query.all()
    addons = Addon.query.filter_by(active=True).order_by(Addon.name).all()
    return render_template('bookings/edit.html', booking=booking, halls=halls, available_addons=[a.to_dict() for a in addons])

@bookings_bp.route('/<int:booking_id>/delete', methods=['POST'])
@login_required
def delete_booking(booking_id):
    """Delete booking"""
    booking = Booking.query.get_or_404(booking_id)
    
    # Check permission
    if not current_user.is_admin() and booking.created_by != current_user.id:
        flash('ليس لديك صلاحيات لحذف هذا الحجز', 'error')
        return redirect(url_for('bookings.list_bookings'))
    
    customer_name = booking.customer_name
    db.session.delete(booking)
    db.session.commit()
    
    # Log action
    log = StaffLog(
        user_id=current_user.id,
        action='delete_booking',
        details=f'حذف الحجز: {customer_name}'
    )
    db.session.add(log)
    db.session.commit()
    
    flash(f'تم حذف الحجز بنجاح', 'success')
    return redirect(url_for('bookings.list_bookings'))

@bookings_bp.route('/export/csv', methods=['GET'])
@login_required
def export_csv():
    """Export bookings to CSV"""
    query = Booking.query
    
    if not current_user.is_admin():
        query = query.filter_by(created_by=current_user.id)
    
    bookings = query.order_by(Booking.date.desc()).all()
    
    # Create CSV
    si = StringIO()
    writer = csv.writer(si)
    writer.writerow(['اسم العميل', 'رقم الموبايل', 'رقم الموبايل البديل', 'القاعة', 'التاريخ', 'الوقت (من-إلى)', 'العربون', 'السعر', 'الإضافات', 'الحالة', 'الملاحظات'])
    
    for booking in bookings:
        hall_name = booking.hall.hall_name if booking.hall else 'غير محدد'
        addons_str = ''
        if booking.addons:
            addons_list = [f"{a.get('name', '')} - {a.get('price', 0)} ج.م" for a in booking.addons]
            addons_str = ' | '.join(addons_list)
        writer.writerow([
            booking.customer_name,
            booking.phone,
            booking.phone_alt or '',
            hall_name,
            booking.date.strftime('%Y-%m-%d'),
            f'{booking.start_time} - {booking.end_time}',
            booking.deposit,
            booking.total_price,
            addons_str,
            booking.status,
            booking.notes or ''
        ])
    
    output = si.getvalue()
    si.close()
    
    from flask import Response
    return Response(
        output,
        mimetype="text/csv",
        headers={"Content-disposition": "attachment;filename=bookings.csv"}
    )

@bookings_bp.route('/export/excel', methods=['GET'])
@login_required
def export_excel():
    """Export bookings to Excel"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        flash('مكتبة Excel غير مثبتة. يرجى تثبيت openpyxl', 'error')
        return redirect(url_for('bookings.list_bookings'))
    
    query = Booking.query
    
    if not current_user.is_admin():
        query = query.filter_by(created_by=current_user.id)
    
    bookings = query.order_by(Booking.date.desc()).all()
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'الحجوزات'
    
    # Headers
    headers = ['اسم العميل', 'رقم الموبايل', 'رقم الموبايل البديل', 'القاعة', 'التاريخ', 'من الساعة', 'إلى الساعة', 'العربون (ج.م)', 'السعر (ج.م)', 'الإضافات', 'الحالة', 'الملاحظات']
    ws.append(headers)
    
    # Style headers
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=12)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    
    # Add data
    for booking in bookings:
        hall_name = booking.hall.hall_name if booking.hall else 'غير محدد'
        status_ar = {
            'pending': 'قيد الانتظار',
            'approved': 'موافق عليه',
            'completed': 'مكتمل',
            'cancelled': 'ملغى'
        }.get(booking.status, booking.status)
        
        addons_str = ''
        if booking.addons:
            addons_list = [f"{a.get('name', '')} - {a.get('price', 0)} ج.م" for a in booking.addons]
            addons_str = '\n'.join(addons_list)
        
        ws.append([
            booking.customer_name,
            booking.phone,
            booking.phone_alt or '',
            hall_name,
            booking.date.strftime('%d/%m/%Y'),
            booking.start_time,
            booking.end_time,
            booking.deposit,
            booking.total_price,
            addons_str,
            status_ar,
            booking.notes or ''
        ])
    
    # Style data cells
    for row in ws.iter_rows(min_row=2, max_row=len(bookings)+1):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = 10
    ws.column_dimensions['H'].width = 12
    ws.column_dimensions['I'].width = 12
    ws.column_dimensions['J'].width = 20
    ws.column_dimensions['K'].width = 12
    ws.column_dimensions['L'].width = 20
    
    # Save to bytes
    from io import BytesIO
    from flask import Response
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return Response(
        output.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment;filename=bookings.xlsx'}
    )
